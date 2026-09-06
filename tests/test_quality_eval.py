from types import SimpleNamespace

from bankrag.config import Settings
from bankrag.quality_eval import DEFAULT_METRICS, METRICS, run_quality
from bankrag.eval_report import run_workbook


class FakeMetric:
    def __init__(self, key):
        self.key = key
        self.score = None
        self.success = None
        self.reason = ""

    def measure(self, tc):
        self.score = 0.9 if "Ferrari" not in tc.actual_output else 0.2
        self.success = self.score >= 0.7
        self.reason = f"{self.key} judged {tc.input[:20]}"


class FakeSession:
    def __init__(self, text):
        self.text = text

    def ask(self, q, force_skill=None):
        return SimpleNamespace(text=self.text, skill_id=force_skill or "credit-card", language="en", trace={"cost": {"total_usd": 0.001}},
                               tool_calls=[{"type": "mcp_call", "name": "knowledge_base_retrieve"}], retrieval_context=["Infinite: 5 lounge visits"])


def _settings(tmp_path):
    return Settings.load(root=tmp_path) if "root" in Settings.load.__code__.co_varnames else Settings.load()


def test_run_quality_with_stubs(tmp_path):
    answers = iter(["5 lounge visits a year.", "5 lounge visits and a free Ferrari."])
    s = Settings.load()
    cases = [{"q": "lounge visits?", "skill": "credit-card"}, {"q": "again?", "expected_output": "5 visits"}]
    logs = []
    run = run_quality(s, {}, cases, metric_keys=["faithfulness", "contextual_recall", "tool_correctness"], threshold=0.7, judge_model="judge-x",
                      session_factory=lambda: FakeSession(next(answers)), metric_factory=FakeMetric, log=logs.append)
    assert run["set"] == "quality" and run["summary"]["judge_model"] == "judge-x"
    r0, r1 = run["rows"]
    assert r0["pass"] is True and r0["metrics"]["faithfulness"]["score"] == 0.9
    assert r0["metrics"]["contextual_recall"]["success"] is None  # skipped: no expected answer
    assert r1["pass"] is False and r1["metrics"]["contextual_recall"]["score"] == 0.2
    assert run["summary"]["passed"] == 1 and run["summary"]["pass_rate"] == 0.5
    assert run["summary"]["avg_scores"]["faithfulness"] == 0.55 and run["summary"]["avg_scores"]["contextual_recall"] == 0.2
    assert any("judge judge-x" in l for l in logs)
    data = run_workbook(run)
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(data))["Results"]
    head = [c.value for c in ws[1]]
    assert head[4] == "faithfulness score" and head[5] == "faithfulness reason" and ws["B3"].value == "FAIL"


def test_metric_registry_and_defaults():
    assert set(DEFAULT_METRICS) <= set(METRICS)
    assert {g for g, _, _, _ in METRICS.values()} == {"rag", "agentic"}
