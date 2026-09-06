import io

from openpyxl import load_workbook

from bankrag.eval_report import history_workbook, run_workbook


ROUTING = {"id": "abc123", "set": "routing", "started_at": "2026-09-06T10:00:00+00:00", "finished_at": "2026-09-06T10:00:09+00:00",
           "summary": {"questions": 2, "passed": 1, "accuracy": 0.5, "total_cost_usd": 0.0012, "avg_ms": 800},
           "rows": [{"q": "บัตรเครดิตใบไหนดี", "expected": "credit-card", "got": "credit-card", "pass": True, "confidence": 0.93, "ms": 700, "cost_usd": 0.0006, "detail": "keywords"},
                    {"q": "weather", "expected": "offtopic", "got": "general", "pass": False, "confidence": 0.4, "ms": 900, "cost_usd": 0.0006, "detail": ""}]}
COMPARE = {"id": "cmp1", "set": "compare", "started_at": "2026-09-06T11:00:00+00:00", "finished_at": "2026-09-06T11:01:00+00:00",
           "summary": {"skill": "credit-card", "models": ["gpt-4.1-mini", "gpt-5.4-mini"], "questions": 1, "gpt-4.1-mini avg_ms": 3000, "gpt-5.4-mini avg_ms": 5000},
           "rows": [{"q": "lounge?", "by_model": {"gpt-4.1-mini": {"text": "5 visits", "ms": 3000, "input_tokens": 100, "output_tokens": 20, "cost_usd": 0.001, "retrieval_calls": 1},
                                                   "gpt-5.4-mini": {"error": "RateLimit", "ms": 5000, "cost_usd": 0.0}}}]}


def test_run_workbook_routing():
    wb = load_workbook(io.BytesIO(run_workbook(ROUTING)))
    assert wb.sheetnames == ["Summary", "Results"]
    s = wb["Summary"]
    assert s["A1"].value == "Field" and s["B2"].value == "abc123"
    labels = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value for r in range(2, s.max_row + 1)}
    assert labels["Accuracy"] == 0.5 and labels["Questions"] == 2
    res = wb["Results"]
    assert [c.value for c in res[1]][:3] == ["#", "Result", "Question"]
    assert res["B2"].value == "PASS" and res["B3"].value == "FAIL" and res["C2"].value == "บัตรเครดิตใบไหนดี"
    assert res.freeze_panes == "A2" and res.auto_filter.ref.startswith("A1:")


def test_run_workbook_compare_and_history():
    wb = load_workbook(io.BytesIO(run_workbook(COMPARE)))
    res = wb["Results"]
    head = [c.value for c in res[1]]
    assert head[2] == "gpt-4.1-mini answer" and "gpt-5.4-mini cost USD" in head
    assert res["C2"].value == "5 visits" and res["I2"].value == "RateLimit"
    hist = load_workbook(io.BytesIO(history_workbook([ROUTING, COMPARE])))
    assert hist.sheetnames[0] == "Runs" and len(hist.sheetnames) == 3
    runs = hist["Runs"]
    assert runs["A2"].value == "abc123" and runs["G2"].value == 0.5 and runs["K3"].value == "gpt-4.1-mini, gpt-5.4-mini"
