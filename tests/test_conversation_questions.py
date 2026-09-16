import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bankrag import api, sessions as SESS
from bankrag.evals import append_cases, load_cases
from bankrag.models import SessionRecord, Turn


def _seed(settings):
    rec = SessionRecord(id="abcdef123456", title="lounge", created_at="2026-09-06T10:00:00+00:00", updated_at="2026-09-06T10:01:04+00:00", turns=[
        Turn(role="user", text="Infinite lounge visits?", at="2026-09-06T10:00:00+00:00"),
        Turn(role="assistant", text="2 visits a year [Infinite](https://x)", at="2026-09-06T10:00:05+00:00", skill_id="credit-card", language="en", confidence=0.9, trace={"cost": {"total_usd": 0.002}, "timings_ms": {"total": 4200}}),
        Turn(role="user", text="ค่าธรรมเนียมรายปี", at="2026-09-06T10:01:00+00:00"),
        Turn(role="assistant", text="3,000 บาทค่ะ", at="2026-09-06T10:01:04+00:00", skill_id="credit-card", language="th", confidence=0.8),
    ])
    SESS.save_session(settings, rec)
    SESS.save_feedback(settings, rec.id, 3, "down", "wrong fee", "pim")


def test_question_rows_filters_and_items(tmp_path, monkeypatch):
    monkeypatch.setenv("SQLITE_DB_PATH", str(tmp_path / "t.db"))
    s = api.settings
    _seed(s)
    rows = SESS.question_rows(s)
    assert [r["question"] for r in rows] == ["ค่าธรรมเนียมรายปี", "Infinite lounge visits?"]
    assert rows[0]["rating"] == "down" and rows[0]["comment"] == "wrong fee" and rows[1]["cost_usd"] == 0.002 and rows[1]["total_ms"] == 4200
    assert [r["idx"] for r in SESS.question_rows(s, rating="down")] == [2]
    assert [r["idx"] for r in SESS.question_rows(s, q="lounge")] == [0]
    assert SESS.question_rows(s, skill="wealth") == []
    assert [r["idx"] for r in SESS.question_rows(s, items=[("abcdef123456", 0)])] == [0]
    # comment filters: any / none / a dislike reason label; the text search also looks in the comment
    SESS.save_feedback(s, "abcdef123456", 1, "up", "", "pim")
    SESS.save_feedback(s, "abcdef123456", 3, "down", "Wrong information; Too slow — fee is 3,500", "pim")
    assert [r["idx"] for r in SESS.question_rows(s, comment="any")] == [2]
    assert [r["idx"] for r in SESS.question_rows(s, comment="none")] == [0]
    assert [r["idx"] for r in SESS.question_rows(s, comment="Too slow")] == [2]
    assert SESS.question_rows(s, comment="Hard to understand") == []
    assert [r["idx"] for r in SESS.question_rows(s, q="3,500")] == [2]
    assert SESS.question_rows(s, comment="any")[0]["comment"] == "Wrong information; Too slow — fee is 3,500"


def test_export_and_append_routes(tmp_path, monkeypatch):
    monkeypatch.setenv("SQLITE_DB_PATH", str(tmp_path / "t.db"))
    monkeypatch.setattr(api.settings, "evals_dir", tmp_path / "evals")
    _seed(api.settings)
    c = TestClient(api.app)
    auth = {"Authorization": "Basic " + __import__("base64").b64encode(f"x:{api.settings.studio_password}".encode()).decode()}
    r = c.post("/conversations/export.xlsx", json={"items": [{"session_id": "abcdef123456", "idx": 2}, {"session_id": "abcdef123456", "idx": 0}]}, headers=auth)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Questions"]
    assert ws["D2"].value == "ค่าธรรมเนียมรายปี" and ws["D3"].value == "Infinite lounge visits?" and ws["I2"].value == "down" and ws["C1"].value == "Asked by"
    r = c.post("/evals/routing/append", json={"cases": [{"q": "Infinite lounge visits?", "skill": "credit-card"}, {"q": "infinite LOUNGE visits?", "skill": "credit-card"}]}, headers=auth)
    assert r.json() == {"added": 1, "skipped": 1, "total": 1}
    assert load_cases(api.settings, "routing") == [{"q": "Infinite lounge visits?", "skill": "credit-card"}]
    assert append_cases(api.settings, "quality", [{"q": "ค่าธรรมเนียมรายปี", "skill": "credit-card"}])["added"] == 1
    assert c.post("/conversations/export.xlsx", json={"items": []}, headers=auth).status_code == 400
