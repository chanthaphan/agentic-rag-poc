import io

from openpyxl import load_workbook

from bankrag.evals import cases_workbook, parse_cases_file


def test_parse_xlsx_roundtrip_all_sets():
    rag = [{"q": "Infinite lounge?", "expect": ["2 visits", "Priority Pass"], "skill": "credit-card"}, {"q": "fee?", "expect": [], "skill": None, "require_source": False}]
    data = cases_workbook("rag", rag)
    ws = load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == ["question", "skill", "expected substrings", "needs source"] and ws["D3"].value == "no"
    parsed = parse_cases_file("rag", data, "x.xlsx")
    assert parsed[0] == {"q": "Infinite lounge?", "expect": ["2 visits", "Priority Pass"], "skill": "credit-card"}
    assert parsed[1]["require_source"] is False and parsed[1]["skill"] is None
    routing = parse_cases_file("routing", cases_workbook("routing", [{"q": "weather?", "skill": "offtopic"}]), "r.xlsx")
    assert routing == [{"q": "weather?", "skill": "offtopic"}]
    quality = parse_cases_file("quality", cases_workbook("quality", [{"q": "q1", "skill": "wealth", "expected_output": "answer"}, {"q": "q2"}]), "q.xlsx")
    assert quality == [{"q": "q1", "skill": "wealth", "expected_output": "answer"}, {"q": "q2"}]


def test_parse_csv_headers_and_headerless():
    csv = "\ufeffQuestion,Expected skill\nบัตรเครดิตใบไหนดี,credit-card\n,ignored\nweather,offtopic\n".encode("utf-8")
    assert parse_cases_file("routing", csv, "list.csv") == [{"q": "บัตรเครดิตใบไหนดี", "skill": "credit-card"}, {"q": "weather", "skill": "offtopic"}]
    plain = "first question\nsecond question\n".encode("utf-8")
    assert [c["q"] for c in parse_cases_file("quality", plain, "notes.csv")] == ["first question", "second question"]
    pipe = "question,expected substrings\nfee?,\"3,000|waived\"\n".encode("utf-8")
    assert parse_cases_file("rag", pipe, "a.csv")[0]["expect"] == ["3,000", "waived"]
