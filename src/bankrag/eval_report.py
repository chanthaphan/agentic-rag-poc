"""Excel reports for eval runs (openpyxl). One workbook per run, or a history workbook with one sheet per run."""
from __future__ import annotations

import io
import re
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="0064FF")
HEAD_FONT = Font(bold=True, color="FFFFFF")
PASS_FILL = PatternFill("solid", fgColor="E4F7E8")
FAIL_FILL = PatternFill("solid", fgColor="FDE8EC")
SUMMARY_LABELS = {"judge_model": "Judge model", "threshold": "Threshold", "metrics": "Metrics", "avg_scores": "Average scores", "questions": "Questions", "passed": "Passed", "accuracy": "Accuracy", "pass_rate": "Pass rate", "total_cost_usd": "Total cost (USD)", "avg_ms": "Avg latency (ms)", "skill": "Skill", "models": "Models"}


def _title(s: str) -> str:
    s = re.sub(r"[\[\]:*?/\\]", "-", s)
    return s[:31] or "sheet"


def _header(ws, cols: Iterable[str], row: int = 1) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = f"A{row + 1}"


def _widths(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _fmt_summary_value(k: str, v):
    if isinstance(v, list):
        return ", ".join(map(str, v))
    if isinstance(v, dict):
        return ", ".join(f"{a}={b:.2f}" if isinstance(b, (int, float)) else f"{a}={b}" for a, b in v.items())
    return v


def _summary_sheet(ws, run: dict) -> None:
    ws.append(["Field", "Value"])
    _header(ws, ["Field", "Value"])
    meta = [("Run id", run.get("id", "")), ("Set", run.get("set", "")), ("Started", run.get("started_at", "")), ("Finished", run.get("finished_at", ""))]
    for k, v in meta:
        ws.append([k, v])
    for k, v in (run.get("summary") or {}).items():
        ws.append([SUMMARY_LABELS.get(k, k), _fmt_summary_value(k, v)])
        if k in ("accuracy", "pass_rate"):
            ws.cell(row=ws.max_row, column=2).number_format = "0.0%"
        elif "cost" in k and isinstance(v, (int, float)):
            ws.cell(row=ws.max_row, column=2).number_format = "$0.0000"
    _widths(ws, {1: 24, 2: 60})


def _results_sheet(ws, run: dict) -> None:
    rows = run.get("rows") or []
    if run.get("set") == "compare":
        models = (run.get("summary") or {}).get("models") or sorted({m for r in rows for m in (r.get("by_model") or {})})
        cols = ["#", "Question"]
        for m in models:
            cols += [f"{m} answer", f"{m} ms", f"{m} in tokens", f"{m} out tokens", f"{m} cost USD", f"{m} retrieval calls"]
        ws.append(cols)
        _header(ws, cols)
        for i, r in enumerate(rows, 1):
            line = [i, r.get("q", "")]
            for m in models:
                x = (r.get("by_model") or {}).get(m) or {}
                line += [x.get("text") or x.get("error") or "", x.get("ms"), x.get("input_tokens"), x.get("output_tokens"), x.get("cost_usd"), x.get("retrieval_calls")]
            ws.append(line)
            for c in range(3, 3 + 6 * len(models), 6):
                ws.cell(row=ws.max_row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
                ws.cell(row=ws.max_row, column=c + 4).number_format = "$0.0000"
            ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        widths = {1: 5, 2: 40}
        for j in range(len(models)):
            base = 3 + 6 * j
            widths.update({base: 60, base + 1: 9, base + 2: 11, base + 3: 11, base + 4: 11, base + 5: 10})
        _widths(ws, widths)
        return
    if run.get("set") == "quality":
        keys = (run.get("summary") or {}).get("metrics") or sorted({k for r in rows for k in (r.get("metrics") or {})})
        cols = ["#", "Result", "Question", "Skill"] + [c for k in keys for c in (f"{k} score", f"{k} reason")] + ["Answer latency ms", "Judge ms", "Cost USD", "Answer", "Expected answer", "Context chunks", "Tools called"]
        ws.append(cols)
        _header(ws, cols)
        for i, r in enumerate(rows, 1):
            line = [i, "PASS" if r.get("pass") else "FAIL", r.get("q", ""), r.get("got") or r.get("skill", "")]
            for k in keys:
                s = (r.get("metrics") or {}).get(k) or {}
                line += [s.get("score"), s.get("reason", "")]
            line += [r.get("ms"), r.get("judge_ms"), r.get("cost_usd"), r.get("answer", ""), r.get("expected_output", ""), r.get("context_chunks"), ", ".join(r.get("tools_called") or [])]
            ws.append(line)
            row = ws.max_row
            ws.cell(row=row, column=2).fill = PASS_FILL if r.get("pass") else FAIL_FILL
            for j, k in enumerate(keys):
                c = ws.cell(row=row, column=5 + 2 * j)
                c.number_format = "0.00"
                s = (r.get("metrics") or {}).get(k) or {}
                if s.get("success") is not None:
                    c.fill = PASS_FILL if s.get("success") else FAIL_FILL
                ws.cell(row=row, column=6 + 2 * j).alignment = Alignment(wrap_text=True, vertical="top")
            base = 5 + 2 * len(keys)
            ws.cell(row=row, column=base + 2).number_format = "$0.0000"
            for c in (3, base + 3, base + 4):
                ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
        widths = {1: 5, 2: 8, 3: 40, 4: 12}
        for j in range(len(keys)):
            widths[5 + 2 * j] = 9
            widths[6 + 2 * j] = 50
        base = 5 + 2 * len(keys)
        widths.update({base: 12, base + 1: 10, base + 2: 10, base + 3: 60, base + 4: 40, base + 5: 9, base + 6: 20})
        _widths(ws, widths)
        return
    cols = ["#", "Result", "Question", "Expected", "Got", "Confidence", "Latency ms", "Cost USD", "Detail"]
    ws.append(cols)
    _header(ws, cols)
    for i, r in enumerate(rows, 1):
        ws.append([i, "PASS" if r.get("pass") else "FAIL", r.get("q", ""), r.get("expected", ""), r.get("got", ""), r.get("confidence"), r.get("ms"), r.get("cost_usd"), r.get("detail", "")])
        row = ws.max_row
        ws.cell(row=row, column=2).fill = PASS_FILL if r.get("pass") else FAIL_FILL
        ws.cell(row=row, column=6).number_format = "0%"
        ws.cell(row=row, column=8).number_format = "$0.0000"
        for c in (3, 9):
            ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
    _widths(ws, {1: 5, 2: 8, 3: 48, 4: 18, 5: 18, 6: 11, 7: 11, 8: 10, 9: 60})


def run_workbook(run: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _summary_sheet(ws, run)
    _results_sheet(wb.create_sheet("Results"), run)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def history_workbook(runs: list[dict], max_sheets: int = 30) -> bytes:
    """Sheet 'Runs' (one row per run) followed by a results sheet per run (newest first, capped)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Runs"
    cols = ["Run id", "Set", "Started", "Finished", "Questions", "Passed", "Score", "Total cost USD", "Avg latency ms", "Skill", "Models"]
    ws.append(cols)
    _header(ws, cols)
    for r in runs:
        s = r.get("summary") or {}
        score = s.get("accuracy", s.get("pass_rate"))
        ws.append([r.get("id"), r.get("set"), r.get("started_at"), r.get("finished_at"), s.get("questions"), s.get("passed"), score, s.get("total_cost_usd"), s.get("avg_ms"), s.get("skill"), ", ".join(s.get("models") or []) or None])
        ws.cell(row=ws.max_row, column=7).number_format = "0.0%"
        ws.cell(row=ws.max_row, column=8).number_format = "$0.0000"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
    _widths(ws, {1: 12, 2: 10, 3: 26, 4: 26, 5: 10, 6: 8, 7: 8, 8: 14, 9: 14, 10: 14, 11: 30})
    for r in runs[:max_sheets]:
        if r.get("rows") is not None:
            _results_sheet(wb.create_sheet(_title(f"{r.get('set')}-{r.get('id')}")), r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def questions_workbook(rows: list[dict]) -> bytes:
    """Selected customer questions with the answer, skill, feedback and cost, one row each."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    cols = ["#", "Asked at", "Asked by", "Question", "Answer", "Skill", "Language", "Confidence", "Rating", "Comment", "Tester", "Cost USD", "Latency ms", "Retrieved docs", "Input tokens", "Output tokens", "Source", "Session", "Session id", "Turn"]
    ws.append(cols)
    _header(ws, cols)
    for i, r in enumerate(rows, 1):
        ws.append([i, r.get("at"), r.get("user"), r.get("question"), r.get("answer"), r.get("skill_id"), r.get("language"), r.get("confidence"), r.get("rating"), r.get("comment"), r.get("tester"),
                   r.get("cost_usd"), r.get("total_ms"), r.get("retrieved_docs"), r.get("input_tokens"), r.get("output_tokens"), r.get("source"), r.get("session_title"), r.get("session_id"), r.get("idx")])
        row = ws.max_row
        for c in (4, 5, 10):
            ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=8).number_format = "0%"
        ws.cell(row=row, column=12).number_format = "$0.0000"
        if r.get("rating") == "up":
            ws.cell(row=row, column=9).fill = PASS_FILL
        elif r.get("rating") == "down":
            ws.cell(row=row, column=9).fill = FAIL_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
    _widths(ws, {1: 5, 2: 22, 3: 18, 4: 45, 5: 70, 6: 12, 7: 9, 8: 10, 9: 8, 10: 30, 11: 12, 12: 10, 13: 10, 14: 9, 15: 10, 16: 10, 17: 8, 18: 30, 19: 14, 20: 6})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DISLIKE_LABELS = ["Wrong information", "Did not answer the question", "Not enough detail", "Hard to understand", "Wrong language or tone", "Too slow"]


def split_comment(comment: str) -> tuple[str, str]:
    """The chat UIs store a dislike as "Label; Label — note": returns (reasons joined by ", ", note)."""
    head, _, rest = (comment or "").partition(" — ")
    labels = [x for x in head.split("; ") if x in DISLIKE_LABELS]
    if not labels:
        return "", comment or ""
    return ", ".join(labels), rest


def chatlog_workbook(rows: list[dict], who: str = "") -> bytes:
    """A person's own chat log: every question with its answer and feedback ("Chat log"), and the rated ones again
    on a "Feedback" sheet with the dislike reasons and the note in their own columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Chat log"
    cols = ["#", "Date", "Conversation", "Question", "Answer", "Topic", "Language", "Rating", "Reasons", "Comment"]
    ws.append(cols)
    _header(ws, cols)
    fb = wb.create_sheet("Feedback")
    fcols = ["#", "Date", "Conversation", "Question", "Answer", "Topic", "Rating", "Reasons", "Comment"]
    fb.append(fcols)
    _header(fb, fcols)
    n_fb = 0
    for i, r in enumerate(sorted(rows, key=lambda x: x.get("at") or ""), 1):
        reasons, note = split_comment(r.get("comment") or "")
        rating = {"up": "👍 helpful", "down": "👎 not helpful"}.get(r.get("rating") or "", "")
        ws.append([i, r.get("at"), r.get("session_title"), r.get("question"), r.get("answer"), r.get("skill_id"), r.get("language"), rating, reasons, note])
        row = ws.max_row
        for c in (4, 5, 10):
            ws.cell(row=row, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        if r.get("rating") == "up":
            ws.cell(row=row, column=8).fill = PASS_FILL
        elif r.get("rating") == "down":
            ws.cell(row=row, column=8).fill = FAIL_FILL
        if r.get("rating") or note or reasons:
            n_fb += 1
            fb.append([n_fb, r.get("at"), r.get("session_title"), r.get("question"), r.get("answer"), r.get("skill_id"), rating, reasons, note])
            frow = fb.max_row
            for c in (4, 5, 9):
                fb.cell(row=frow, column=c).alignment = Alignment(wrap_text=True, vertical="top")
            fb.cell(row=frow, column=7).fill = PASS_FILL if r.get("rating") == "up" else FAIL_FILL
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
    fb.auto_filter.ref = f"A1:{get_column_letter(len(fcols))}{max(1, fb.max_row)}"
    _widths(ws, {1: 5, 2: 22, 3: 28, 4: 45, 5: 70, 6: 14, 7: 9, 8: 14, 9: 30, 10: 40})
    _widths(fb, {1: 5, 2: 22, 3: 28, 4: 45, 5: 70, 6: 14, 7: 14, 8: 30, 9: 40})
    if who:
        ws.cell(row=1, column=len(cols) + 2, value=f"Exported by {who}").font = Font(italic=True, color="888888")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

