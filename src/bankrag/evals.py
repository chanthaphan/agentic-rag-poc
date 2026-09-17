"""Evaluation runs (routing accuracy, grounded answers, model comparison) shared by the CLI and Studio."""
from __future__ import annotations

import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Optional

import yaml

from .config import Settings
from .ingest.progress import report as progress
from .models import SkillSpec

Log = Callable[[str], None]
SETS = {"routing": "routing_questions.yaml", "rag": "rag_questions.yaml", "quality": "quality_questions.yaml"}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def eval_path(settings: Settings, set_name: str) -> Path:
    if set_name not in SETS:
        raise ValueError("unknown eval set")
    return settings.evals_dir / SETS[set_name]


def load_cases(settings: Settings, set_name: str) -> list[dict]:
    p = eval_path(settings, set_name)
    return list(yaml.safe_load(p.read_text(encoding="utf-8")) or []) if p.exists() else []


def save_cases(settings: Settings, set_name: str, cases: list[dict]) -> list[dict]:
    clean = []
    for c in cases:
        q = str(c.get("q", "")).strip()
        if not q:
            continue
        if set_name == "routing":
            clean.append({"q": q, "skill": str(c.get("skill", "")).strip()})
        elif set_name == "quality":
            row = {"q": q}
            if c.get("skill"):
                row["skill"] = str(c["skill"])
            if str(c.get("expected_output") or "").strip():
                row["expected_output"] = str(c["expected_output"]).strip()
            clean.append(row)
        else:
            row = {"q": q, "expect": [str(x) for x in (c.get("expect") or []) if str(x).strip()]}
            if c.get("skill"):
                row["skill"] = str(c["skill"])
            if c.get("require_source") is False:
                row["require_source"] = False
            clean.append(row)
    p = eval_path(settings, set_name)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(yaml.safe_dump(clean, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return clean


def append_cases(settings: Settings, set_name: str, cases: list[dict]) -> dict:
    """Add questions to a set, skipping any whose text already exists (case-insensitive). Returns counts."""
    existing = load_cases(settings, set_name)
    seen = {str(c.get("q", "")).strip().lower() for c in existing}
    added = 0
    for c in cases:
        q = str(c.get("q", "")).strip()
        if not q or q.lower() in seen:
            continue
        existing.append(c)
        seen.add(q.lower())
        added += 1
    saved = save_cases(settings, set_name, existing)
    return {"added": added, "skipped": len(cases) - added, "total": len(saved)}


COLUMN_ALIASES = {
    "q": ("q", "question", "questions", "คำถาม", "customer question", "input"),
    "skill": ("skill", "expected skill", "expected_skill", "skill (optional)", "skill id", "skill_id", "routed skill"),
    "expect": ("expect", "expected", "expected substrings", "expected substrings (comma)", "expected text", "must contain"),
    "require_source": ("require_source", "needs source", "requires source", "needs_source", "source required"),
    "expected_output": ("expected_output", "expected answer", "expected output", "expected answer (optional)", "reference answer", "ground truth"),
}
SET_COLUMNS = {"routing": ["q", "skill"], "rag": ["q", "skill", "expect", "require_source"], "quality": ["q", "skill", "expected_output"]}
COLUMN_TITLES = {"q": "question", "skill": "skill", "expect": "expected substrings", "require_source": "needs source", "expected_output": "expected answer"}


def _canon(header: str) -> Optional[str]:
    h = str(header or "").strip().lower()
    for key, names in COLUMN_ALIASES.items():
        if h in names:
            return key
    return None


def _truthy(v) -> bool:
    return str(v).strip().lower() in ("1", "true", "yes", "y", "x", "✓", "ใช่")


def parse_cases_file(set_name: str, data: bytes, filename: str = "") -> list[dict]:
    """Read questions from an .xlsx (first sheet) or .csv. The header row names the columns (see COLUMN_ALIASES);
    a file without a recognised header is read as one question per row in the first column."""
    if set_name not in SETS:
        raise ValueError("unknown eval set")
    rows: list[list] = []
    if filename.lower().endswith(".csv") or (not filename.lower().endswith((".xlsx", ".xlsm")) and not data.startswith(b"PK")):
        import csv
        import io

        rows = [r for r in csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))]
    else:
        import io

        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header = [_canon(c) for c in rows[0]]
    if "q" in header:
        body = rows[1:]
    else:
        header = ["q"] + [None] * (len(rows[0]) - 1)
        body = rows
    idx = {k: header.index(k) for k in ("q", "skill", "expect", "require_source", "expected_output") if k in header}
    out: list[dict] = []
    for r in body:
        get = lambda k: str(r[idx[k]]).strip() if k in idx and idx[k] < len(r) and r[idx[k]] is not None else ""  # noqa: E731
        q = get("q")
        if not q:
            continue
        if set_name == "routing":
            out.append({"q": q, "skill": get("skill")})
        elif set_name == "quality":
            c = {"q": q}
            if get("skill"):
                c["skill"] = get("skill")
            if get("expected_output"):
                c["expected_output"] = get("expected_output")
            out.append(c)
        else:
            raw = get("expect")
            expect = [x.strip() for x in (raw.split("|") if "|" in raw else raw.split(",")) if x.strip()]
            c = {"q": q, "expect": expect, "skill": get("skill") or None}
            if "require_source" in idx and get("require_source") and not _truthy(get("require_source")):
                c["require_source"] = False
            out.append(c)
    return out


def cases_workbook(set_name: str, cases: list[dict]) -> bytes:
    """The question set as an .xlsx with the same columns the upload accepts (so it doubles as a template)."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cols = SET_COLUMNS[set_name]
    wb = Workbook()
    ws = wb.active
    ws.title = {"routing": "Routing", "rag": "Grounded", "quality": "Quality"}[set_name]
    ws.append([COLUMN_TITLES[c] for c in cols])
    for cell in ws[1]:
        cell.fill, cell.font = PatternFill("solid", fgColor="0064FF"), Font(bold=True, color="FFFFFF")
    ws.freeze_panes = "A2"
    for c in cases:
        line = []
        for k in cols:
            v = c.get(k)
            if k == "expect":
                v = ", ".join(v or [])
            elif k == "require_source":
                v = "no" if v is False else "yes"
            line.append("" if v is None else v)
        ws.append(line)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    widths = {"q": 60, "skill": 16, "expect": 40, "require_source": 12, "expected_output": 60}
    for i, k in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[k]
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, ws.max_row)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def new_run(set_name: str) -> dict:
    return {"id": uuid.uuid4().hex[:10], "set": set_name, "started_at": _now(), "rows": [], "summary": {}}


def run_routing(settings: Settings, skills: dict[str, SkillSpec], cases: list[dict], *, llm=None, log: Log = print) -> dict:
    from . import router as R
    from .pricing import load_pricing, usage_cost

    if llm is None:
        from .llm import chat_model

        llm = chat_model(settings, settings.router_model)
    pricing = load_pricing(settings)
    run = new_run("routing")
    ok = 0
    progress(log, "questions", 0, len(cases), message="", passed=0, failed=0)
    for i, c in enumerate(cases):
        progress(log, "questions", i, len(cases), message=c["q"][:80], passed=ok, failed=i - ok)
        t0 = time.perf_counter()
        d = R.route(llm, c["q"], [], None, skills)
        ms = int((time.perf_counter() - t0) * 1000)
        cost = usage_cost(pricing, settings.router_model, d.usage)["total_usd"] if d.usage else 0.0
        hit = d.skill_id == c["skill"]
        ok += hit
        run["rows"].append({"q": c["q"], "expected": c["skill"], "got": d.skill_id, "pass": hit, "confidence": d.confidence, "ms": ms, "cost_usd": cost, "detail": d.reason[:160]})
        log(f"{'ok ' if hit else 'BAD'} expected={c['skill']} got={d.skill_id} conf={d.confidence:.2f} {c['q'][:60]}")
    n = len(cases)
    progress(log, "questions", n, n, message="", passed=ok, failed=n - ok)
    run["summary"] = {"questions": n, "passed": ok, "accuracy": (ok / n) if n else 0.0, "total_cost_usd": sum(r["cost_usd"] for r in run["rows"]), "avg_ms": int(sum(r["ms"] for r in run["rows"]) / n) if n else 0}
    run["finished_at"] = _now()
    return run


def run_rag(settings: Settings, skills: dict[str, SkillSpec], cases: list[dict], *, session_factory=None, log: Log = print) -> dict:
    from .chat import ChatSession

    factory = session_factory or (lambda: ChatSession(settings, skills))
    run = new_run("rag")
    passed = 0
    progress(log, "questions", 0, len(cases), message="", passed=0, failed=0)
    for i, c in enumerate(cases):
        progress(log, "questions", i, len(cases), message=c["q"][:80], passed=passed, failed=i - passed)
        session = factory()
        t0 = time.perf_counter()
        try:
            ans = session.ask(c["q"], force_skill=c.get("skill"))
        except Exception as e:  # noqa: BLE001
            run["rows"].append({"q": c["q"], "expected": ", ".join(c.get("expect", [])), "got": f"ERROR {type(e).__name__}", "pass": False, "ms": int((time.perf_counter() - t0) * 1000), "cost_usd": 0.0, "detail": str(e)[:160]})
            continue
        ms = int((time.perf_counter() - t0) * 1000)
        expect = c.get("expect", [])
        matched = [e for e in expect if e in ans.text]
        has_src = bool(ans.citations or ans.references)
        good = len(matched) == len(expect) and (has_src or not c.get("require_source", True))
        passed += good
        run["rows"].append({"q": c["q"], "expected": ", ".join(expect) or "(any)", "got": ans.skill_id, "pass": good, "confidence": ans.confidence, "ms": ms,
                            "cost_usd": (ans.trace.get("cost") or {}).get("total_usd", 0.0), "detail": f"matched {len(matched)}/{len(expect)}; sources={'yes' if has_src else 'no'} · {ans.text[:140]}"})
        log(f"{'ok ' if good else 'BAD'} skill={ans.skill_id} matched={len(matched)}/{len(expect)} sources={has_src} {c['q'][:60]}")
        try:
            session.reset()
        except Exception:  # noqa: BLE001
            pass
    n = len(cases)
    progress(log, "questions", n, n, message="", passed=passed, failed=n - passed)
    run["summary"] = {"questions": n, "passed": passed, "pass_rate": (passed / n) if n else 0.0, "total_cost_usd": sum(r["cost_usd"] for r in run["rows"]), "avg_ms": int(sum(r["ms"] for r in run["rows"]) / n) if n else 0}
    run["finished_at"] = _now()
    return run


def run_compare(settings: Settings, skills: dict[str, SkillSpec], base_body: str, skill_id: str, models: list[str], questions: list[str], *, log: Log = print,
                session_factory=None) -> dict:
    """The same skill definition with the model swapped, each question on a fresh in-memory thread; nothing to clean up."""
    from langgraph.checkpoint.memory import InMemorySaver

    from .chat import ChatSession
    from .llm import chat_model

    spec = skills[skill_id]
    run = new_run("compare")
    run["summary"] = {"skill": skill_id, "models": models, "questions": len(questions)}
    rows = [{"q": q, "by_model": {}} for q in questions]
    total = len(models) * len(questions)
    done = 0
    progress(log, "questions", 0, total, message="")
    for m in models:
        skills_m = {**skills, skill_id: spec.model_copy(update={"model": m})}
        for row in rows:
            progress(log, "questions", done, total, message=f"{m}: {row['q'][:60]}")
            done += 1
            t0 = time.perf_counter()
            try:
                session = session_factory(m) if session_factory else ChatSession(
                    settings, skills_m, llm_factory=lambda _model, m=m: chat_model(settings, m), checkpointer=InMemorySaver(), base_body=base_body)
                session.kb_owners = {**session.kb_owners}  # per-session copy; the model swap does not move the knowledge base
                # the live (unpublished) definition carries the swapped model: bypass the registry for this one skill
                from .sync import desired_definition

                session._definition_for = lambda sp, _d=desired_definition, _o=session.kb_owners: (_d(settings, sp, base_body, _o.get(sp.id, sp)), "compare")
                ans = session.ask(row["q"], force_skill=skill_id, with_sources=False)
                usage = (ans.trace.get("usage") or {}).get("agent") or {}
                cost = (ans.trace.get("cost") or {})
                row["by_model"][m] = {"text": ans.text, "ms": int((time.perf_counter() - t0) * 1000), "input_tokens": usage.get("input_tokens", 0), "output_tokens": usage.get("output_tokens", 0),
                                      "cost_usd": float((cost.get("agent") or {}).get("total_usd", cost.get("total_usd", 0.0)) or 0.0),
                                      "retrieval_calls": (ans.trace.get("retrieval") or {}).get("calls", 0), "documents": (ans.trace.get("retrieval") or {}).get("documents", 0)}
                log(f"[{m}] {row['q'][:50]} -> {row['by_model'][m]['ms']} ms, ${row['by_model'][m]['cost_usd']:.4f}")
            except Exception as e:  # noqa: BLE001
                row["by_model"][m] = {"error": f"{type(e).__name__}: {str(e)[:200]}", "ms": int((time.perf_counter() - t0) * 1000), "cost_usd": 0.0}
                log(f"[{m}] ERROR {row['by_model'][m]['error']}")
    run["rows"] = rows
    for m in models:
        vals = [r["by_model"].get(m, {}) for r in rows]
        run["summary"][f"{m} avg_ms"] = int(sum(v.get("ms", 0) for v in vals) / max(1, len(vals)))
        run["summary"][f"{m} cost_usd"] = sum(v.get("cost_usd", 0.0) for v in vals)
    run["finished_at"] = _now()
    return run
